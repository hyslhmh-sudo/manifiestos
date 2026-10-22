from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FUENTE = 'Plus Jakarta Sans'
AZUL = RGBColor(0x0D, 0x31, 0x3F); PAT = RGBColor(0x58, 0x94, 0xA7); GRIS = RGBColor(0x3D, 0x4A, 0x52); GRIS2 = RGBColor(0x6B, 0x77, 0x80)

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.left_margin = s.right_margin = Mm(22); s.top_margin = Mm(24); s.bottom_margin = Mm(22)

st = d.styles['Normal']; st.font.name = FUENTE; st.font.size = Pt(10.5); st.font.color.rgb = GRIS
st.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
pf = st.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(4); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for n, size in (('Heading 1', 15), ('Heading 2', 12)):
    h = d.styles[n]; h.font.name = FUENTE; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = AZUL
    h.element.rPr.rFonts.set(qn('w:asciiTheme'), ''); h.element.rPr.rFonts.set(qn('w:ascii'), FUENTE); h.element.rPr.rFonts.set(qn('w:hAnsi'), FUENTE)
    h.paragraph_format.space_before = Pt(14 if size > 13 else 10); h.paragraph_format.space_after = Pt(4); h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def campo(run, instr):
    for t, v in (('begin', None), (None, instr), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); run._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = v; run._r.append(e)

# Encabezado: logo HMH
hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.add_run().add_picture('/home/claude/desvios/hmh-000.png', width=Mm(48))
r = hp.add_run('\tServicio de Higiene y Seguridad Laboral'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS2
# Pie: logo HySL centrado + «Página X de Y»
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run().add_picture('/home/claude/manifiestos/desvios/logo-hysl.png', width=Mm(120))
fp2 = s.footer.add_paragraph(); fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for txt, fld in (('Página ', None), (None, 'PAGE'), (' de ', None), (None, 'NUMPAGES')):
    r = fp2.add_run(txt or ''); r.font.size = Pt(8); r.font.color.rgb = GRIS2
    if fld: campo(r, fld)

def P(texto, bold_ini=None, size=None, color=None, align=None, after=None):
    p = d.add_paragraph()
    if bold_ini:
        r = p.add_run(bold_ini); r.bold = True
    r = p.add_run(texto)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def L(items, num=False):
    for it in items:
        p = d.add_paragraph(style='List Number' if num else 'List Bullet')
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.3
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)
        for r in p.runs: r.font.name = FUENTE; r.font.color.rgb = GRIS
def tabla(filas, anchos):
    t = d.add_table(rows=0, cols=len(anchos)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, f in enumerate(filas):
        cells = t.add_row().cells
        for j, v in enumerate(f):
            c = cells[j]; c.width = Mm(anchos[j]); p = c.paragraphs[0]; p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.1; p.paragraph_format.space_after = Pt(2)
            r = p.add_run(v); r.font.size = Pt(9.5); r.font.name = FUENTE
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else GRIS; r.bold = i == 0
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), '0D313F'); c._tc.get_or_add_tcPr().append(sh)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- Portada breve
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT; t.paragraph_format.space_before = Pt(6)
r = t.add_run('Guía de uso'); r.font.size = Pt(11); r.font.color.rgb = PAT; r.bold = True
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run('Registro de desvíos en el recorrido de retiro de residuos'); r.font.size = Pt(20); r.bold = True; r.font.color.rgb = AZUL
t.paragraph_format.line_spacing = 1.15
P('Para los operarios de Star Union que hacen el traslado interno de residuos en el Hospital María Humphreys (HMH).', size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
P('Versión 0 · octubre de 2026 · Servicio de Higiene y Seguridad Laboral HZT / HMH', size=9, color=GRIS2, align=WD_ALIGN_PARAGRAPH.LEFT, after=10)

d.add_heading('1. Para qué sirve', level=1)
P('En cada recorrido ustedes ven antes que nadie lo que está mal en la gestión de residuos: bolsas en el piso, bolsas del color equivocado, agujas o vidrios sueltos, contenedores llenos o depósitos desordenados. Este registro sirve para que eso llegue en el momento a Higiene y Seguridad, con lugar, hora y foto, y se pueda corregir en el servicio que lo genera.')
P('No es para buscar culpables. Cada registro es una oportunidad de mejora y ayuda a cuidar a quien manipula los residuos después.')

d.add_heading('2. Cómo entrar', level=1)
L([('Escaneando el QR ', 'pegado en el punto del recorrido: la página abre con el lugar ya elegido.'),
   ('Desde el enlace ', 'hyslhmh-sudo.github.io/manifiestos/desvios (sirve para cualquier lugar).'),
   ('Consejo: ', 'en el celular, tocá los tres puntitos del navegador y «Agregar a la pantalla de inicio». Queda como una aplicación más.')])

d.add_heading('3. Paso a paso', level=1)
L([('Quién carga. ', 'La primera vez elegí tu nombre y escribí los 3 últimos números de tu DNI. El teléfono te recuerda para la próxima. Si usás un celular compartido, tocá «No soy yo» y elegí tu nombre.'),
   ('Dónde. ', 'Si entraste por el QR, el lugar ya está elegido. Si no, elegí planta baja o planta alta y tocá el lugar. Si querés, agregá más precisión (por ejemplo, «tacho del fondo, al lado del ascensor»).'),
   ('Qué viste. ', 'Marcá una o varias opciones. Las de cortopunzantes se ven en rojo porque llegan como urgentes.'),
   ('Qué residuo es ', 'y si había residuos mezclados.'),
   ('Contalo en pocas palabras. ', 'Por ejemplo: «bolsa roja abierta en el piso del depósito, tacho vacío al lado».'),
   ('Fotos. ', 'Podés sacar hasta 5. Que no aparezcan pacientes ni personas.'),
   ('Gravedad y lesión. ', 'Marcá qué tan grave te parece y si alguien se lastimó.'),
   ('Enviar. ', 'Aparece el número de registro (por ejemplo, DR-0042). La fecha y la hora se guardan solas.')], num=True)

d.add_heading('4. Qué marcar en «¿Qué viste?»', level=1)
tabla([('Opción', 'Ejemplo'),
       ('Residuo fuera de contenedor', 'Bolsas rojas o negras en el piso; tacho vacío y bolsa al lado.'),
       ('Bolsa incorrecta / Segregación incorrecta', 'Bandejas de comida en bolsa roja; residuo con sangre en bolsa negra.'),
       ('Bolsa abierta / sin cerrar · Falta de bolsa', 'Bolsa roja sin cerrar; tacho sin bolsa.'),
       ('Contenedor lleno / desbordado · Contenedor inadecuado', 'Tacho rebalsado; cajas de cartón usadas como contenedor.'),
       ('Cortopunzante mal dispuesto', 'Agujas, ampollas o vidrios sueltos en bolsa o caja (llega como URGENTE).'),
       ('Depósito sin condiciones adecuadas', 'Depósito sucio, desordenado o con cosas que no son residuos.'),
       ('Residuo químico sin identificación · Falta de rotulado', 'Bidón o frasco sin etiqueta; recipiente sin rótulo.'),
       ('Otro', 'Cualquier otra cosa: contalo en la descripción.')], [62, 104])

d.add_heading('5. Si alguien se lastimó', level=1)
P('Marcá «Sí» en «¿Alguien se lastimó?» y contá qué pasó. La carga llega como URGENTE a Higiene y Seguridad. Además:')
L(['Avisá en el momento a tu supervisor de Star Union.', 'Andá a la guardia para la atención y para que quede registrado.', 'Si es un pinchazo o corte con material usado, no esperes: la atención en las primeras horas es importante.'])

d.add_heading('6. Sin señal', level=1)
P('Si no hay señal, la carga queda guardada en el teléfono y se manda sola cuando vuelve la conexión. Arriba a la derecha aparece «Sin señal · 1 sin enviar». No la cargues de nuevo.')

d.add_heading('7. Cuidados durante el registro', level=1)
L([('Mirá sin tocar. ', 'Nunca abras una bolsa cerrada para ver qué tiene.'),
   ('No levantes cortopunzantes con la mano. ', 'Usá los elementos de protección y el descartador.'),
   ('Primero la seguridad, después la foto. ', 'Si algo es peligroso (derrame, agujas sueltas), avisá en el momento a tu supervisor.')])

d.add_heading('8. Preguntas frecuentes', level=1)
tabla([('Pregunta', 'Respuesta'),
       ('Mi nombre no aparece.', 'Avisale a tu supervisor o a Higiene y Seguridad para que te agreguen.'),
       ('Me dice que el DNI no coincide.', 'Revisá los 3 últimos números. Si sigue sin andar, avisá a Higiene y Seguridad.'),
       ('Me equivoqué en una carga.', 'Hacé una carga nueva aclarando en «Algo más» qué corregís, o avisá a Higiene y Seguridad.'),
       ('La página no abre.', 'Probá de nuevo con señal. Higiene y Seguridad tiene una dirección de respaldo.'),
       ('¿Quién ve lo que cargo?', 'Higiene y Seguridad y su secretaría. Se usa para mejorar la gestión de residuos de los servicios.')], [62, 104])

d.add_heading('9. Marco', level=1)
p = P('Este registro forma parte de la gestión interna de residuos del HMH y de las tareas de prevención del Servicio de Higiene y Seguridad Laboral (Ley N.º 19.587 y Decreto N.º 351/79; Ley N.º 24.051 de residuos peligrosos y Ley XI N.º 35 de la Provincia del Chubut).', size=9.5, color=GRIS2)

d.save('/home/claude/desvios/Guia_registro_desvios_HMH_v0.docx'); print('ok')
