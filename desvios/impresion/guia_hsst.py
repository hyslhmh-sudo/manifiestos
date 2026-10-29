from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FUENTE = 'Plus Jakarta Sans'
AZUL = RGBColor(0x0D, 0x31, 0x3F); PAT = RGBColor(0x58, 0x94, 0xA7); GRIS = RGBColor(0x3D, 0x4A, 0x52); GRIS2 = RGBColor(0x6B, 0x77, 0x80)

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.left_margin = s.right_margin = Mm(22); s.top_margin = Mm(24); s.bottom_margin = Mm(22)

st = d.styles['Normal']; st.font.name = FUENTE; st.font.size = Pt(10.5); st.font.color.rgb = GRIS
st.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
pf = st.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(4); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for n, size in (('Heading 1', 15), ('Heading 2', 12)):
    h = d.styles[n]; h.font.name = FUENTE; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = AZUL
    h.element.rPr.rFonts.set(qn('w:asciiTheme'), ''); h.element.rPr.rFonts.set(qn('w:ascii'), FUENTE); h.element.rPr.rFonts.set(qn('w:hAnsi'), FUENTE)
    h.paragraph_format.space_before = Pt(14 if size > 13 else 10); h.paragraph_format.space_after = Pt(4); h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def campo(run, instr):
    for t, v in (('begin', None), (None, instr), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); run._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = v; run._r.append(e)

# Encabezado: logo HMH
hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.add_run().add_picture('/home/claude/hyslhmh-sudo.github.io/starh/enc-hsst.png', width=Mm(26))
r = hp.add_run('\tServicio de Higiene y Seguridad Laboral'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS2
# Pie: logo HySL centrado + «Página X de Y»
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run().add_picture('/home/claude/manifiestos/desvios/logo-hysl.png', width=Mm(120))
fp2 = s.footer.add_paragraph(); fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for txt, fld in (('Página ', None), (None, 'PAGE'), (' de ', None), (None, 'NUMPAGES')):
    r = fp2.add_run(txt or ''); r.font.size = Pt(8); r.font.color.rgb = GRIS2
    if fld: campo(r, fld)

def P(texto, bold_ini=None, size=None, color=None, align=None, after=None):
    p = d.add_paragraph()
    if bold_ini:
        r = p.add_run(bold_ini); r.bold = True
    r = p.add_run(texto)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def L(items, num=False):
    for it in items:
        p = d.add_paragraph(style='List Number' if num else 'List Bullet')
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.3
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)
        for r in p.runs: r.font.name = FUENTE; r.font.color.rgb = GRIS
def tabla(filas, anchos):
    t = d.add_table(rows=0, cols=len(anchos)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, f in enumerate(filas):
        cells = t.add_row().cells
        for j, v in enumerate(f):
            c = cells[j]; c.width = Mm(anchos[j]); p = c.paragraphs[0]; p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.1; p.paragraph_format.space_after = Pt(2)
            r = p.add_run(v); r.font.size = Pt(9.5); r.font.name = FUENTE
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else GRIS; r.bold = i == 0
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), '0D313F'); c._tc.get_or_add_tcPr().append(sh)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- Portada breve
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT; t.paragraph_format.space_before = Pt(6)
r = t.add_run('Guía de uso'); r.font.size = Pt(11); r.font.color.rgb = PAT; r.bold = True
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run('Registro de desvíos en la gestión de residuos'); r.font.size = Pt(20); r.bold = True; r.font.color.rgb = AZUL
t.paragraph_format.line_spacing = 1.15
P('Para el referente del Hospital Subzonal Santa Teresita (HSST, Rawson): internación, servicios y CAPS.', size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
P('Servicio de Higiene y Seguridad Laboral · Complementa y no reemplaza al Servicio de Higiene y Seguridad del establecimiento.', size=9, color=GRIS2, align=WD_ALIGN_PARAGRAPH.LEFT, after=10)

d.add_heading('1. Para qué sirve', level=1)
P('Para dejar registrado lo que se ve mal con los residuos en la internación, en los servicios y en los CAPS: bolsas en el piso o del color equivocado, cortopunzantes sueltos, contenedores llenos o depósitos desordenados. No es para buscar culpables: cada registro es una oportunidad de mejora.')
P('No se mandan correos. Lo cargado queda en el sistema: Higiene y Seguridad lo revisa y el referente lo ve en «Mi establecimiento».')

d.add_heading('2. Cómo entrar', level=1)
L([('Desde «Mi establecimiento» ', '› 1 · Residuos › «Cargar un desvío».'),
   ('Escaneando el QR ', 'pegado en cada servicio o CAPS: la página abre con el lugar ya elegido.'),
   ('Desde el enlace ', 'hyslhmh-sudo.github.io/manifiestos/desvios/?e=HSST'),
   ('Consejo: ', 'en el celular, tocá los tres puntitos del navegador y «Agregar a la pantalla de inicio».')])

d.add_heading('3. Paso a paso', level=1)
L([('Quién carga. ', 'La primera vez elegí tu nombre y escribí los 3 últimos números de tu DNI. El teléfono te recuerda para la próxima.'),
   ('Dónde. ', 'Si entraste por el QR, el lugar ya está elegido. Si no, elegí Internación, Servicios o CAPS y tocá el lugar. Podés agregar más precisión (por ejemplo, «habitación 4», «depósito del fondo»).'),
   ('Antes se registró en este lugar. ', 'Si Higiene y Seguridad ya tomó una medida por un desvío de ese lugar, la página lo muestra: marcá «Sigue igual» o «Ya está corregido».'),
   ('Qué viste. ', 'Marcá una o varias opciones. Las de cortopunzantes se ven en rojo porque quedan como urgentes.'),
   ('Por qué pasó. ', 'En «Contenedor lleno» o «Residuo fuera de contenedor», mirá todos los contenedores: si no había lugar en ninguno, el problema es que faltan contenedores.'),
   ('Aislamientos. ', 'Todo lo que sale de un aislamiento por enfermedad infectocontagiosa, incluida la comida, va en bolsa roja: no es un error.'),
   ('Fotos. ', 'Hasta 5, sacadas en el momento o elegidas del teléfono. Que no aparezcan pacientes ni personas.'),
   ('Enviar. ', 'Aparece el número de registro. La fecha y la hora se guardan solas.')], num=True)

d.add_heading('4. Ver lo cargado y verificar', level=1)
P('En «Mi establecimiento» › 1 · Residuos › «Mis desvíos» están todos los registros, lo que revisó Higiene y Seguridad y lo que hay que verificar en el lugar. Cuando pases por el lugar, tocá «Corregido» o «Sigue igual».')

d.add_heading('5. Si alguien se lastimó', level=1)
P('Marcá «Sí» en «¿Alguien se lastimó?» y contá qué pasó. Además, la persona tiene que ir a la guardia en el momento; si es un pinchazo o corte con material usado, no hay que esperar.')

d.add_heading('6. Sin señal', level=1)
P('La carga queda guardada en el teléfono y se manda sola cuando vuelve la conexión. No la cargues de nuevo.')

d.add_heading('7. Cuidados', level=1)
L([('Mirá sin tocar. ', 'Nunca abras una bolsa cerrada.'), ('No levantes cortopunzantes con la mano.', ''), ('Primero la seguridad, después la foto.', '')])

d.add_heading('8. Marco', level=1)
P('Este registro forma parte de la gestión interna de residuos del establecimiento y de las tareas de prevención en Higiene y Seguridad (Ley N.º 19.587 y Decreto N.º 351/79; Ley N.º 24.051 de residuos peligrosos y Ley XI N.º 35 de la Provincia del Chubut).')

d.save('/home/claude/manifiestos/desvios/impresion/Guia_registro_desvios_HSST.docx'); print('ok')
